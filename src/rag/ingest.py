"""Ingesta: extracción de texto desde TXT, PDF, DOCX y HTML."""
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".html", ".htm"}


@dataclass
class Document:
    source: str                 # nombre del archivo (se usa para citar)
    text: str
    page: Optional[int] = None  # solo PDF


def _read_txt(path: Path):
    return [Document(path.name, path.read_text(encoding="utf-8", errors="ignore"))]


def _read_pdf(path: Path):
    from pypdf import PdfReader
    docs = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(path.name, text, page=i))
    return docs


def _read_docx(path: Path):
    from docx import Document as DocxDocument
    d = DocxDocument(str(path))
    lines = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        if p.style is not None and p.style.name.lower().startswith(("heading", "título", "titulo")):
            lines.append(f"## {p.text.strip()}")
        else:
            lines.append(p.text.strip())
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [Document(path.name, "\n".join(lines))]


def _read_html(path: Path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    for h in soup.find_all(["h1", "h2", "h3", "h4"]):
        h.string = f"\n## {h.get_text(strip=True)}\n"
    return [Document(path.name, soup.get_text("\n"))]


_READERS = {".txt": _read_txt, ".md": _read_txt, ".pdf": _read_pdf,
            ".docx": _read_docx, ".html": _read_html, ".htm": _read_html}


def load_documents(docs_dir: Path):
    docs = []
    for path in sorted(Path(docs_dir).iterdir()):
        if path.suffix.lower() in SUPPORTED:
            docs.extend(_READERS[path.suffix.lower()](path))
    if not docs:
        raise FileNotFoundError(f"No se encontraron documentos soportados en {docs_dir}")
    return docs
